"""
Document Text Extractor for RADIX JD Analytics.
Extracts clean, readable text from multiple file formats:
PDF (.pdf), Word (.docx, .doc), Plain Text (.txt, .md), Rich Text (.rtf), HTML (.html, .htm), and JSON (.json).
"""

import re
from io import BytesIO
from pathlib import Path
from typing import Optional

try:
    import pymupdf  # PyMuPDF
except ImportError:
    pymupdf = None

try:
    from pypdf import PdfReader
except ImportError:
    PdfReader = None

try:
    from docx import Document
except ImportError:
    Document = None


def extract_text(filename: str, content: bytes) -> str:
    """
    Extracts plain text content from binary file bytes based on file extension.

    Supports:
      - PDF (.pdf) via PyMuPDF or pypdf
      - Word (.docx, .doc) via python-docx or binary string parsing
      - Plain text (.txt, .md, .markdown, .text, .log)
      - Rich text (.rtf)
      - Web / Data (.html, .htm, .json, .csv)
    """
    if not content or len(content) == 0:
        raise ValueError("Uploaded document is empty (0 bytes).")

    extension = Path(filename).suffix.lower()

    if extension == ".pdf":
        text = _extract_pdf(content)

    elif extension == ".docx":
        text = _extract_docx(content)

    elif extension == ".doc":
        text = _extract_legacy_doc(content)

    elif extension in (".txt", ".md", ".markdown", ".text", ".log", ".json", ".csv"):
        text = _extract_plain_text(content)

    elif extension in (".html", ".htm"):
        raw_html = _extract_plain_text(content)
        # Strip script, style, and HTML tags
        clean_html = re.sub(r"<(?:script|style)[^>]*>[\s\S]*?</(?:script|style)>", " ", raw_html, flags=re.I)
        text = re.sub(r"<[^>]+>", " ", clean_html)

    elif extension == ".rtf":
        raw_rtf = _extract_plain_text(content)
        # Basic RTF control word stripping
        text = re.sub(r"\\[a-zA-Z0-9]+(?:\s|-?[0-9]+)?", " ", raw_rtf)
        text = re.sub(r"[{}\\]", " ", text)

    else:
        # Check for binary or unparseable extensions
        if extension in (".bin", ".exe", ".dll", ".so", ".dat", ".zip", ".tar", ".gz", ".7z", ".png", ".jpg", ".jpeg"):
            raise ValueError(f"Binary file type '{extension}' is not supported. Please upload a PDF, Word, or text document.")

        # Default attempt: if unknown extension, check if printable text
        try:
            text = _extract_plain_text(content)
            printable_count = sum(1 for ch in text if ch.isprintable() or ch in "\n\r\t")
            if len(text) > 0 and (printable_count / len(text)) < 0.75:
                raise ValueError("Binary content detected.")
        except Exception:
            raise ValueError(
                f"Unsupported file format '{extension}'. Please upload a PDF, Word (DOCX), or plain text document."
            )

    cleaned = _clean_text(text)
    if not cleaned:
        raise ValueError(
            "No readable text found in document. If this is a scanned PDF, OCR may be required."
        )

    return cleaned


def _extract_pdf(content: bytes) -> str:
    """Extracts text from PDF content using PyMuPDF or pypdf."""
    if pymupdf is not None:
        try:
            with pymupdf.open(stream=content, filetype="pdf") as doc:
                parts = [page.get_text() for page in doc]
                return "\n".join(parts)
        except Exception as exc:
            # Fall back to pypdf if PyMuPDF had an issue
            if PdfReader is not None:
                return _extract_pypdf(content)
            raise ValueError(f"Failed to parse PDF document: {exc}")

    elif PdfReader is not None:
        return _extract_pypdf(content)

    raise ValueError("No PDF extraction library available (neither pymupdf nor pypdf is installed).")


def _extract_pypdf(content: bytes) -> str:
    reader = PdfReader(BytesIO(content))
    return "\n".join(page.extract_text() or "" for page in reader.pages)


def _extract_docx(content: bytes) -> str:
    """Extracts text and table contents from a .docx file."""
    if Document is None:
        raise ValueError("python-docx is not installed to parse .docx documents.")

    try:
        doc = Document(BytesIO(content))
    except Exception as exc:
        raise ValueError(f"Corrupt or invalid DOCX document: {exc}")

    parts = []
    for paragraph in doc.paragraphs:
        if paragraph.text.strip():
            parts.append(paragraph.text.strip())

    for table in doc.tables:
        for row in table.rows:
            row_cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
            if row_cells:
                parts.append(" | ".join(row_cells))

    return "\n".join(parts)


def _extract_legacy_doc(content: bytes) -> str:
    """
    Extracts text from legacy .doc files.
    First tries python-docx in case it's actually an OOXML file with .doc extension,
    then falls back to extracting printable string runs.
    """
    try:
        return _extract_docx(content)
    except Exception:
        pass

    # Extract sequences of ASCII printable characters from binary stream
    text_pieces = re.findall(rb"[\x20-\x7E\r\n\t]{4,}", content)
    decoded = [p.decode("latin1", errors="ignore").strip() for p in text_pieces if len(p.strip()) > 3]
    joined = "\n".join(decoded)
    if len(joined) > 30:
        return joined

    raise ValueError(
        "Legacy .doc binary format cannot be parsed. Please convert the file to modern .docx or export as .pdf."
    )


def _extract_plain_text(content: bytes) -> str:
    """Decodes plain text with resilient encoding fallbacks."""
    for encoding in ("utf-8", "utf-8-sig", "latin-1", "cp1252"):
        try:
            return content.decode(encoding)
        except (UnicodeDecodeError, LookupError):
            continue
    return content.decode("utf-8", errors="ignore")


def _clean_text(raw: str) -> str:
    """Normalizes whitespace and removes null characters."""
    if not raw:
        return ""
    # Strip null characters
    clean = raw.replace("\x00", "")
    # Normalize Windows CRLF to LF
    clean = clean.replace("\r\n", "\n").replace("\r", "\n")
    return clean.strip()
