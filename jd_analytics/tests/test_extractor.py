"""
Unit tests for document text extractor (PDF, DOCX, TXT, MD).
"""

import pytest
from io import BytesIO
from docx import Document

from jd_analytics.extractor import extract_text


def test_extract_plain_text():
    content = b"Looking for a Python Developer with experience in SQL and Docker."
    text = extract_text("sample.txt", content)
    assert "Python Developer" in text
    assert "Docker" in text


def test_extract_markdown():
    content = b"# Job Description\n## Required Skills\n- Python\n- Kubernetes\n- FastAPI"
    text = extract_text("job_spec.md", content)
    assert "Job Description" in text
    assert "Kubernetes" in text


def test_extract_docx():
    doc = Document()
    doc.add_heading("Software Engineer JD", 0)
    doc.add_paragraph("We require extensive background in Java, Spring Boot, and PostgreSQL.")
    bio = BytesIO()
    doc.save(bio)

    text = extract_text("spec.docx", bio.getvalue())
    assert "Software Engineer JD" in text
    assert "Java" in text
    assert "PostgreSQL" in text


def test_extract_empty_file_raises():
    with pytest.raises(ValueError, match="empty"):
        extract_text("empty.txt", b"")


def test_unsupported_extension_with_unparseable_content():
    # A completely arbitrary binary file
    with pytest.raises(ValueError):
        extract_text("unsupported.bin", b"\x00\x01\x02\x03\x04\x00\x00")
