"""
API Integration tests for RADIX Resume Parsing Microservice.
Tests endpoints: /health, /resume-parsing/extract, /resume-parsing/samples, /resume-parsing/{id}, DELETE.
"""

from io import BytesIO
import pytest
from fastapi.testclient import TestClient
from app.main import app
from app.storage import clear_all_resumes


@pytest.fixture(autouse=True)
def clean_storage():
    """Ensure clean storage state before each test."""
    clear_all_resumes()
    yield
    clear_all_resumes()


def test_health_check():
    client = TestClient(app)
    response = client.get("/health")
    assert response.status_code == 200
    data = response.json()
    assert data["status"] == "ok"
    assert data["service"] == "resume_parsing"
    assert data["port"] == 8002


def test_empty_storage_initially():
    """Verify that no mock sample data is seeded by default."""
    client = TestClient(app)
    response = client.get("/resume-parsing/samples")
    assert response.status_code == 200
    assert response.json() == []


def test_extract_from_text_upload():
    client = TestClient(app)
    resume_content = b"""
    Vikram Malhotra
    vikram.m@example.com | +91 91234 56789
    
    Education:
    B.Tech Computer Science, IIT Delhi, 2025, 9.1 CGPA
    
    Skills:
    Python, C++, Data Structures, Algorithms, Docker, PostgreSQL, Linux
    """

    response = client.post(
        "/resume-parsing/extract",
        files={"file": ("vikram_resume.txt", BytesIO(resume_content), "text/plain")},
    )
    assert response.status_code == 200
    data = response.json()
    assert data["candidate_name"] == "Vikram Malhotra"
    assert data["candidate_email"] == "vikram.m@example.com"
    assert data["source_type"] == "resume"
    assert len(data["skills"]) >= 3

    # Check persistence: GET /resume-parsing/samples should now return this resume
    get_res = client.get("/resume-parsing/samples")
    assert get_res.status_code == 200
    stored_list = get_res.json()
    assert len(stored_list) == 1
    assert stored_list[0]["id"] == data["id"]

    # Check GET by ID
    single_res = client.get(f"/resume-parsing/{data['id']}")
    assert single_res.status_code == 200
    assert single_res.json()["candidate_name"] == "Vikram Malhotra"

    # Check DELETE by ID
    del_res = client.delete(f"/resume-parsing/{data['id']}")
    assert del_res.status_code == 200

    # Storage should be empty again
    assert len(client.get("/resume-parsing/samples").json()) == 0


def test_extract_sample_candidate_ananya():
    """Test compatibility with test_end_to_end.py candidate-ananya-rao."""
    client = TestClient(app)
    response = client.post(
        "/resume-parsing/extract",
        data={"sample_candidate_id": "candidate-ananya-rao"}
    )
    assert response.status_code == 200
    data = response.json()
    assert data["candidate_name"] == "Ananya Rao"
    assert data["source_type"] == "resume"
    assert len(data["skills"]) >= 5


def test_extract_docx_upload():
    import docx
    doc = docx.Document()
    doc.add_heading("Priya Sharma", 0)
    doc.add_paragraph("priya.sharma@example.com | +91 99887 76655")
    doc.add_paragraph("B.Tech in CSE, 2024. Skilled in Java, Spring Boot, SQL, and Git.")

    bio = BytesIO()
    doc.save(bio)
    docx_bytes = bio.getvalue()

    client = TestClient(app)
    response = client.post(
        "/resume-parsing/extract",
        files={"file": ("priya_resume.docx", BytesIO(docx_bytes), "application/vnd.openxmlformats-officedocument.wordprocessingml.document")},
    )
    assert response.status_code == 200
    data = response.json()
    assert data["candidate_name"] == "Priya Sharma"
    assert data["candidate_email"] == "priya.sharma@example.com"
    assert len(data["skills"]) >= 3
