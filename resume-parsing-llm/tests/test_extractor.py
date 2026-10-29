"""
Unit tests for resume document text extraction.
"""

from io import BytesIO
import pytest
from app.extractor import extract_text


def test_extract_plain_text():
    sample_text = "John Doe\nSoftware Engineer\nSkills: Python, Go, Docker"
    extracted = extract_text("resume.txt", sample_text.encode("utf-8"))
    assert "John Doe" in extracted
    assert "Python" in extracted


def test_extract_markdown():
    sample_md = "# Candidate: Alice Smith\n## Education\nB.Tech in CSE from IIT Bombay"
    extracted = extract_text("cv.md", sample_md.encode("utf-8"))
    assert "Alice Smith" in extracted
    assert "IIT Bombay" in extracted


def test_extract_docx():
    import docx
    doc = docx.Document()
    doc.add_heading("Karthik Narayan", 0)
    doc.add_paragraph("Data Scientist with PyTorch and SQL expertise.")
    table = doc.add_table(rows=1, cols=2)
    table.rows[0].cells[0].text = "Experience"
    table.rows[0].cells[1].text = "2 Years"

    bio = BytesIO()
    doc.save(bio)
    docx_bytes = bio.getvalue()

    extracted = extract_text("resume.docx", docx_bytes)
    assert "Karthik Narayan" in extracted
    assert "PyTorch" in extracted
    assert "Experience" in extracted


def test_extract_empty_fails():
    with pytest.raises(ValueError, match="empty"):
        extract_text("empty.txt", b"")


def test_unsupported_binary_fails():
    with pytest.raises(ValueError, match="Binary file type"):
        extract_text("binary.exe", b"\x00\x01\x02\x03\x04")
